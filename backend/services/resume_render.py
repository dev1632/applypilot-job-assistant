"""Render a verified draft to DOCX. Names, titles, organizations, and dates always come from the
profile, never from model output, so they cannot drift."""
from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from backend.models.profile import Profile
from backend.models.resume import TailoredDraft

MONTHS = "Jan Feb Mar Apr May Jun Jul Aug Sep Oct Nov Dec".split()


def fmt_month(value: str | None) -> str:
    if not value:
        return "Present"
    year, month = value.split("-")[:2]
    return f"{MONTHS[int(month) - 1]} {year}"


def _heading(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text.upper())
    run.bold = True
    run.font.size = Pt(11)
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(2)


def _entry_line(doc: Document, left: str, right: str) -> None:
    p = doc.add_paragraph()
    p.add_run(left).bold = True
    if right:
        p.add_run(f"  |  {right}")
    p.paragraph_format.space_after = Pt(0)


def render_docx(profile: Profile, draft: TailoredDraft, path: Path, coursework: list[str] | None = None) -> Path:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)
    doc.core_properties.author = profile.name
    doc.core_properties.title = f"{profile.name} Resume"

    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(profile.name)
    run.bold = True
    run.font.size = Pt(18)
    contact = [x for x in (profile.email, profile.phone, *profile.links.values()) if x]
    line = doc.add_paragraph(" | ".join(contact))
    line.alignment = WD_ALIGN_PARAGRAPH.CENTER

    if draft.summary:
        _heading(doc, "Summary")
        doc.add_paragraph(draft.summary.text)

    if profile.education:
        _heading(doc, "Education")
        today = date.today().strftime("%Y-%m")
        for i, edu in enumerate(profile.education):
            grad = fmt_month(edu.expected_graduation)
            when = f"Expected {grad}" if edu.expected_graduation > today else grad
            _entry_line(doc, f"{edu.degree} in {edu.field}, {edu.school}", when)
            courses = coursework if (i == 0 and coursework is not None) else edu.coursework
            if courses:
                doc.add_paragraph(f"Relevant coursework: {', '.join(courses)}")

    if draft.skills:
        _heading(doc, "Skills")
        doc.add_paragraph(", ".join(draft.skills))

    exp_by_id = {f"exp:{e.id}": e for e in profile.experience}
    if draft.experiences:
        _heading(doc, "Experience")
        for entry in draft.experiences:
            exp = exp_by_id[entry.entry_id]
            _entry_line(doc, f"{exp.title}, {exp.organization}", f"{fmt_month(exp.start)} – {fmt_month(exp.end)}")
            for b in entry.bullets:
                doc.add_paragraph(b.text, style="List Bullet")

    proj_by_id = {f"proj:{p.id}": p for p in profile.projects}
    if draft.projects:
        _heading(doc, "Projects")
        for entry in draft.projects:
            proj = proj_by_id[entry.entry_id]
            _entry_line(doc, proj.name, ", ".join(proj.technologies))
            for b in entry.bullets:
                doc.add_paragraph(b.text, style="List Bullet")

    path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(path)
    path.chmod(0o600)
    return path


def draft_to_text(draft: TailoredDraft) -> str:
    """Plain-text view used for previews and keyword checks."""
    parts = [draft.summary.text] if draft.summary else []
    for e in draft.experiences + draft.projects:
        parts.append(e.entry_id)
        parts += [f"• {b.text}" for b in e.bullets]
    parts.append("Skills: " + ", ".join(draft.skills))
    return "\n".join(parts)
