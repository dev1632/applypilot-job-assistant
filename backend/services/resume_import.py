"""Import an existing resume into a *draft* profile the user reviews before anything is saved.

The uploaded file is stored privately and only parsed for text (never executed or rendered). The
model may only copy what the resume says; code then checks every bullet, number, skill, and
technology against the extracted text and flags anything that isn't there. Nothing reaches the
verified profile until the user picks it in the review screen.
"""
from __future__ import annotations

import io
import json
import re
import uuid
from datetime import datetime, timezone
from pathlib import Path

from pydantic import BaseModel, Field

from backend.config import get_settings
from backend.models.profile import Education, Experience, Profile, Project
from backend.services.llm import LLMClient
from backend.services.scoring import _mentions, canon

MAX_BYTES = 5 * 1024 * 1024
SIGNATURES = {".docx": b"PK\x03\x04", ".pdf": b"%PDF-"}
DATE = re.compile(r"^\d{4}-(0[1-9]|1[0-2])$")


class ImportError_(ValueError):
    pass


# ---- Upload and text extraction ----------------------------------------------------------------

def store_upload(data: bytes, filename: str) -> Path:
    suffix = Path(filename or "").suffix.lower()
    if suffix not in (".docx", ".pdf", ".txt"):
        raise ImportError_("Upload a .docx, .pdf, or .txt resume")
    if not data:
        raise ImportError_("The file is empty")
    if len(data) > MAX_BYTES:
        raise ImportError_("The file is larger than 5 MB")
    if suffix in SIGNATURES and not data.startswith(SIGNATURES[suffix]):
        raise ImportError_(f"This doesn't look like a real {suffix} file")
    folder = get_settings().data_dir / "master_resume"
    folder.mkdir(parents=True, exist_ok=True)
    folder.chmod(0o700)
    stem = re.sub(r"[^A-Za-z0-9_-]+", "_", Path(filename).stem)[:60] or "resume"
    path = folder / f"{datetime.now(timezone.utc):%Y%m%dT%H%M%S}_{stem}{suffix}"
    path.write_bytes(data)
    path.chmod(0o600)
    return path


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".txt":
        try:
            text = path.read_text(encoding="utf-8")
        except UnicodeDecodeError:
            raise ImportError_("The text file isn't UTF-8") from None
    elif suffix == ".docx":
        from docx import Document

        doc = Document(str(path))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
        text = "\n".join(parts)
    else:
        from pypdf import PdfReader

        try:
            reader = PdfReader(io.BytesIO(path.read_bytes()))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:10])
        except Exception as exc:
            raise ImportError_(f"Couldn't read this PDF ({exc.__class__.__name__})") from None
    text = text.replace(" ", " ").strip()
    if len(text) < 50:
        raise ImportError_("Couldn't find text in this file. Scanned PDFs aren't supported; try a .docx")
    return text[:60000]


# ---- Model extraction --------------------------------------------------------------------------

class ImpLink(BaseModel):
    label: str = Field(description="linkedin, github, website, or portfolio")
    url: str


class ImpEducation(BaseModel):
    school: str
    degree: str
    field: str
    graduation: str | None = Field(description="YYYY-MM if the resume gives month and year, else null")
    coursework: list[str]


class ImpExperience(BaseModel):
    title: str
    organization: str
    start: str | None = Field(description="YYYY-MM if stated with month, else null")
    end: str | None = Field(description="YYYY-MM if stated with month, else null (also null if current)")
    current: bool
    bullets: list[str] = Field(description="Copied word for word from the resume, without the bullet symbol")
    technologies: list[str] = Field(description="Only technologies the resume names for this role")


class ImpProject(BaseModel):
    name: str
    description: str = Field(description="Copied word for word: the project's first line or sentence")
    technologies: list[str]
    outcomes: list[str] = Field(description="Remaining bullets for the project, copied word for word")


class ImportedProfile(BaseModel):
    name: str
    email: str | None
    phone: str | None
    links: list[ImpLink]
    skills: list[str]
    education: list[ImpEducation]
    experience: list[ImpExperience]
    projects: list[ImpProject]


IMPORT_INSTRUCTIONS = """You convert a resume into structured data. COPY, don't write:
- Bullets, descriptions, and outcomes must be copied word for word from the resume (drop only the
  bullet symbol). Never rephrase, merge, shorten, or add words.
- Dates: use YYYY-MM only when the resume states the month and year; otherwise null. Never guess.
- Education graduation is the END of the date range (e.g. "Jan 2025 – Dec 2026" -> 2026-12), or the
  single date if only one is given.
- PDF text wraps long bullets onto several lines: join the lines of one bullet back together.
- Technologies and skills: only ones the resume actually names.
- Leave out anything the resume doesn't contain. Don't infer work authorization or other personal data."""


def llm_extract(llm: LLMClient, text: str) -> ImportedProfile:
    return llm.structured(instructions=IMPORT_INSTRUCTIONS, task="Extract this resume.", document=text,
                          schema=ImportedProfile)


EMAIL = re.compile(r"[\w.+-]+@[\w-]+\.[\w.-]+")
PHONE = re.compile(r"(?:\+?\d{1,2}[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}")
URL = re.compile(r"(?:https?://)?(?:www\.)?(linkedin\.com/in/[\w-]+|github\.com/[\w-]+)", re.I)


def rule_extract(text: str) -> ImportedProfile:
    """Without OpenAI: contact details and a Skills line only. Experience is left for the user."""
    first_line = next((ln.strip() for ln in text.splitlines() if ln.strip()), "")
    links = [ImpLink(label="linkedin" if "linkedin" in m.group(1).lower() else "github", url=f"https://{m.group(1)}")
             for m in URL.finditer(text)]
    skills: list[str] = []
    m = re.search(r"^\s*(technical )?skills?\s*[:\n](.{0,600})", text, re.I | re.M | re.S)
    if m:
        block = re.split(r"\n\s*\n|\n[A-Z][A-Z ]{3,}\n", m.group(2))[0]
        skills = [s.strip(" •·-*\t") for s in re.split(r"[,|;•·\n]", block) if 1 < len(s.strip(" •·-*\t")) <= 40]
        skills = [re.sub(r"^[A-Za-z ]+:\s*", "", s) for s in skills]  # "Languages: Python" -> "Python"
    return ImportedProfile(
        name=first_line if len(first_line.split()) <= 5 and not EMAIL.search(first_line) else "",
        email=(EMAIL.search(text) or [None])[0], phone=(PHONE.search(text) or [None])[0],
        links=list({link.url: link for link in links}.values()), skills=list(dict.fromkeys(s for s in skills if s)),
        education=[], experience=[], projects=[])


# ---- Verification ------------------------------------------------------------------------------

def _norm(s: str) -> str:
    s = s.replace("–", "-").replace("—", "-").replace("’", "'").replace("“", '"').replace("”", '"')
    return " ".join(re.sub(r"^[\s•·▪●◦*\-–]+", "", s).split()).lower()


def _in_text(fragment: str, text_norm: str) -> bool:
    return bool(fragment.strip()) and _norm(fragment) in text_norm


def _numbers(s: str) -> set[str]:
    return set(re.findall(r"\d[\d,.]*", s))


MONTHS = {m: i for i, m in enumerate("jan feb mar apr may jun jul aug sep oct nov dec".split(), 1)}
MONTH_YEAR = re.compile(r"\b(jan|feb|mar|apr|may|jun|jul|aug|sep|oct|nov|dec)[a-z]*\.?\s+(\d{4})\b", re.I)


def graduation_from_text(school: str, text: str) -> str | None:
    """The last month-year on the school's line (the end of its date range), if the resume has one."""
    for line in text.splitlines():
        if school and _norm(school) in _norm(line):
            dates = MONTH_YEAR.findall(line)
            if dates:
                mon, year = dates[-1]
                return f"{year}-{MONTHS[mon[:3].lower()]:02d}"
    return None


def join_wrapped(parts: list[str]) -> list[str]:
    """Rejoin bullets that PDF line-wrapping split in two: a piece that doesn't end a sentence followed
    by one that starts in lowercase (or ends with a dangling 'and'/comma) is one bullet."""
    out: list[str] = []
    for part in (p.strip() for p in parts):
        if not part:
            continue
        if out and (not re.search(r"[.!?:;)]\s*$", out[-1]) and
                    (part[:1].islower() or re.search(r"(,|\b(and|or|with|to|of|for|the|by|in))\s*$", out[-1], re.I))):
            out[-1] = f"{out[-1]} {part}"
        else:
            out.append(part)
    return out


BULLET_START = re.compile(r"^\s*[•·▪●◦*\-–]\s+")


def text_bullets(text: str) -> list[str]:
    """The resume's own bullets, with wrapped lines joined: a line starting with a bullet symbol begins
    a new one; following lines that don't start a new bullet or a heading continue it."""
    units: list[str] = []
    current: str | None = None
    for line in text.splitlines():
        stripped = line.strip()
        if not stripped:
            current = None
            continue
        if BULLET_START.match(stripped):
            current = BULLET_START.sub("", stripped)
            units.append(current)
        elif current is not None and not re.search(r"[.!?]\s*$", units[-1]) and not stripped.isupper():
            units[-1] = current = f"{units[-1]} {stripped}"
        else:
            current = None
    return units


def complete_from_text(fragment: str, units: list[str]) -> str:
    """If a bullet stops mid-sentence and the resume has a longer bullet starting with it, use that."""
    if re.search(r"[.!?]\s*$", fragment):
        return fragment
    f = _norm(fragment)
    matches = [u for u in units if _norm(u).startswith(f) and len(_norm(u)) > len(f)]
    return matches[0] if len(matches) == 1 else fragment


def repair(draft: ImportedProfile, text: str) -> tuple[ImportedProfile, list[str]]:
    """Fix known extraction slips against the resume text itself. Returns notes on what changed."""
    d = draft.model_copy(deep=True)
    notes: list[str] = []
    for e in d.education:
        expected = graduation_from_text(e.school, text)
        if expected and e.graduation != expected:
            notes.append(f"{e.school}: graduation set to {expected}, the end of the dates on your resume "
                         f"(the reader had {e.graduation or 'none'})")
            e.graduation = expected
    units = text_bullets(text)
    for x in d.experience:
        x.bullets = list(dict.fromkeys(complete_from_text(b, units) for b in join_wrapped(x.bullets)))
    for p in d.projects:
        pieces = [complete_from_text(b, units) for b in join_wrapped([p.description, *p.outcomes])]
        pieces = [b for i, b in enumerate(pieces) if not any(_norm(b) in _norm(o) for o in pieces[:i])]
        if pieces:
            p.description, p.outcomes = pieces[0], pieces[1:]
    return d, notes


def annotate(draft: ImportedProfile, text: str, method: str) -> dict:
    tn = _norm(text)
    # PDFs often break lines mid-sentence; compare against a version with line breaks joined too.
    tn_joined = _norm(re.sub(r"-?\n", " ", text))
    found = lambda s: _in_text(s, tn) or _in_text(s, tn_joined)  # noqa: E731
    numbers_in_text = _numbers(text)

    def bullet(s: str) -> dict:
        flags = []
        if not found(s):
            flags.append("Not found word for word in your resume")
        extra = _numbers(s) - numbers_in_text
        if extra:
            flags.append(f"Contains numbers not in your resume: {', '.join(sorted(extra))}")
        return {"text": s.strip(), "verbatim": not flags, "flags": flags}

    def tech(values: list[str]) -> list[dict]:
        return [{"value": v, "found": _mentions(canon(v), tn) or found(v)} for v in dict.fromkeys(values) if v.strip()]

    contact = []
    for field, value in (("name", draft.name), ("email", draft.email), ("phone", draft.phone)):
        if value:
            contact.append({"field": field, "value": value.strip(), "found": found(value) or
                            (field == "phone" and re.sub(r"\D", "", value) in re.sub(r"\D", "", text))})
    for link in draft.links:
        contact.append({"field": f"links.{link.label.lower()}", "value": link.url,
                        "found": found(re.sub(r"^https?://(www\.)?", "", link.url).rstrip("/"))})

    def date_flags(start, end, current) -> list[str]:
        out = []
        if not start or not DATE.match(start):
            out.append("Start month missing; add it before importing")
        if end and not DATE.match(end):
            out.append("End date isn't YYYY-MM")
        if start and end and DATE.match(start) and DATE.match(end) and end < start:
            out.append("End is before start")
        return out

    warnings = []
    if method == "rules":
        warnings.append("No OpenAI key: only contact details and skills were read. Add experience and projects "
                        "in the profile editor, or set OPENAI_API_KEY and import again.")
    return {
        "method": method,
        "contact": contact,
        "skills": tech(draft.skills),
        "education": [{"id": f"edu{i}", **e.model_dump(),
                       "flags": [] if found(e.school) else ["School name not found in your resume"]}
                      for i, e in enumerate(draft.education)],
        "experience": [{"id": f"exp{i}", "title": e.title, "organization": e.organization, "start": e.start,
                        "end": None if e.current else e.end, "current": e.current,
                        "technologies": tech(e.technologies), "bullets": [bullet(b) for b in e.bullets],
                        "flags": date_flags(e.start, None if e.current else e.end, e.current)
                        + ([] if found(e.organization) else ["Organization not found in your resume"])}
                       for i, e in enumerate(draft.experience)],
        "projects": [{"id": f"proj{i}", "name": p.name, "description": bullet(p.description),
                      "technologies": tech(p.technologies), "outcomes": [bullet(o) for o in p.outcomes],
                      "flags": [] if found(p.name) else ["Project name not found in your resume"]}
                     for i, p in enumerate(draft.projects)],
        "warnings": warnings,
    }


# ---- Drafts on disk ----------------------------------------------------------------------------

def _drafts_dir() -> Path:
    d = get_settings().data_dir / "imports"
    d.mkdir(parents=True, exist_ok=True)
    d.chmod(0o700)
    return d


def create_import(data: bytes, filename: str, llm: LLMClient | None) -> dict:
    path = store_upload(data, filename)
    text = extract_text(path)
    method, draft, failure = "rules", None, None
    if llm is not None:
        try:
            draft, method = llm_extract(llm, text), "openai"
        except Exception as exc:  # fall back rather than fail the upload
            failure = f"OpenAI extraction failed ({exc.__class__.__name__}); used the basic reader"
    if draft is None:
        draft = rule_extract(text)
    draft, repairs = repair(draft, text)
    record = {"id": uuid.uuid4().hex[:16], "created_at": datetime.now(timezone.utc).isoformat(),
              "file": path.name, "chars": len(text), **annotate(draft, text, method)}
    record["warnings"].extend(repairs)
    if failure:
        record["warnings"].insert(0, failure)
    out = _drafts_dir() / f"{record['id']}.json"
    out.write_text(json.dumps(record, indent=2))
    out.chmod(0o600)
    return record


def load_import(import_id: str) -> dict:
    if not re.fullmatch(r"[0-9a-f]{16}", import_id):
        raise ImportError_("Unknown import")
    path = _drafts_dir() / f"{import_id}.json"
    if not path.exists():
        raise ImportError_("Unknown import")
    return json.loads(path.read_text())


# ---- Merge -------------------------------------------------------------------------------------

class ExperiencePick(BaseModel):
    id: str
    start: str
    end: str | None = None
    bullets: list[int] = []  # indexes of bullets to keep


class ProjectPick(BaseModel):
    id: str
    outcomes: list[int] = []


class ImportSelection(BaseModel):
    contact: list[str] = []  # field names to take from the resume (overwrites)
    skills: list[str] = []
    education: list[str] = []
    experience: list[ExperiencePick] = []
    projects: list[ProjectPick] = []


def _unique_id(base: str, taken: set[str]) -> str:
    base = re.sub(r"[^a-z0-9]+", "-", base.lower()).strip("-")[:30] or "item"
    candidate, n = base, 2
    while candidate in taken:
        candidate, n = f"{base}-{n}", n + 1
    taken.add(candidate)
    return candidate


def apply_import(profile: Profile, record: dict, pick: ImportSelection) -> tuple[Profile, list[str]]:
    """Merge only what the user picked. Existing profile entries are never modified."""
    p = profile.model_copy(deep=True)
    log: list[str] = []
    contact = {c["field"]: c["value"] for c in record["contact"]}
    for field in pick.contact:
        if field not in contact:
            continue
        if field.startswith("links."):
            p.links[field.split(".", 1)[1]] = contact[field]
        elif field in ("name", "email", "phone"):
            setattr(p, field, contact[field])
        log.append(f"Set {field}")

    have = {canon(s) for s in p.skills}
    allowed = {s["value"] for s in record["skills"]}
    for s in pick.skills:
        if s in allowed and canon(s) not in have:
            p.skills.append(s)
            have.add(canon(s))
            log.append(f"Added skill {s}")

    by_id = {e["id"]: e for e in record["education"]}
    for eid in pick.education:
        e = by_id.get(eid)
        if not e:
            continue
        existing = next((x for x in p.education if x.school.lower() == e["school"].lower()
                         and x.degree.lower() == e["degree"].lower()), None)
        if existing:
            differs = (e["graduation"] or "") != existing.expected_graduation
            log.append(f"Skipped {e['degree']} at {e['school']}: already in profile"
                       + (" with different dates; edit it in the profile editor to update" if differs else ""))
            continue
        p.education.append(Education(school=e["school"], degree=e["degree"], field=e["field"],
                                     expected_graduation=e["graduation"] or "", coursework=e["coursework"]))
        log.append(f"Added education: {e['degree']}, {e['school']}")

    taken = {x.id for x in p.experience}
    by_id = {e["id"]: e for e in record["experience"]}
    for ep in pick.experience:
        e = by_id.get(ep.id)
        if not e:
            continue
        if not DATE.match(ep.start) or (ep.end and not DATE.match(ep.end)) or (ep.end and ep.end < ep.start):
            raise ImportError_(f"{e['title']} at {e['organization']}: dates must be YYYY-MM with end after start")
        bullets = [e["bullets"][i]["text"] for i in ep.bullets if 0 <= i < len(e["bullets"])]
        existing = next((x for x in p.experience if x.organization.lower() == e["organization"].lower()
                         and x.title.lower() == e["title"].lower()), None)
        if existing:
            differs = (existing.bullets, existing.start, existing.end) != (bullets, ep.start, ep.end)
            log.append(f"Skipped {e['title']} at {e['organization']}: already in profile"
                       + (" with different content; edit it in the profile editor to update" if differs else ""))
            continue
        techs = [t["value"] for t in e["technologies"] if t["found"]]
        p.experience.append(Experience(id=_unique_id(e["organization"], taken), title=e["title"],
                                       organization=e["organization"], start=ep.start, end=ep.end,
                                       bullets=bullets, technologies=techs))
        log.append(f"Added {e['title']} at {e['organization']} ({len(bullets)} bullets)")

    taken = {x.id for x in p.projects}
    by_id = {x["id"]: x for x in record["projects"]}
    for pp in pick.projects:
        x = by_id.get(pp.id)
        if not x:
            continue
        outcomes = [x["outcomes"][i]["text"] for i in pp.outcomes if 0 <= i < len(x["outcomes"])]
        existing = next((y for y in p.projects if y.name.lower() == x["name"].lower()), None)
        if existing:
            differs = (existing.description, existing.outcomes) != (x["description"]["text"], outcomes)
            log.append(f"Skipped project {x['name']}: already in profile"
                       + (" with different content; edit it in the profile editor to update" if differs else ""))
            continue
        p.projects.append(Project(id=_unique_id(x["name"], taken), name=x["name"], description=x["description"]["text"],
                                  technologies=[t["value"] for t in x["technologies"] if t["found"]], outcomes=outcomes))
        log.append(f"Added project {x['name']}")
    return p, log
